"""Plain text from uploaded files, split into sections (a PDF page each, otherwise the whole file)."""
import io
import re
from dataclasses import dataclass
from pathlib import PurePath

SUPPORTED_EXTENSIONS = (".txt", ".md", ".markdown", ".pdf", ".docx")


class DocumentError(ValueError):
    """An upload that cannot be indexed; status_code is the HTTP status to answer with."""

    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None


def clean_text(text: str) -> str:
    """Normalise line endings and whitespace; PostgreSQL text cannot hold NUL characters."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extension_of(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def extract_sections(filename: str, data: bytes) -> list[Section]:
    """Sections of non-empty text, or DocumentError for unsupported or unreadable files."""
    extension = extension_of(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"Unsupported file type {extension or '(none)'}; upload one of: {', '.join(SUPPORTED_EXTENSIONS)}.",
            status_code=415,
        )
    if extension == ".pdf":
        sections = _pdf_sections(data)
    elif extension == ".docx":
        sections = [Section(_docx_text(data))]
    else:
        sections = [Section(_decode(data))]

    sections = [Section(clean_text(s.text), s.page) for s in sections]
    sections = [s for s in sections if s.text]
    if not sections:
        raise DocumentError("The file contains no extractable text (a scanned PDF needs OCR first).")
    return sections


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")  # common for text saved on Windows


def _pdf_sections(data: bytes) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("The PDF is password-protected.")
        return [Section(page.extract_text() or "", number) for number, page in enumerate(reader.pages, 1)]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentError(f"The PDF could not be read: {exc}") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for a bad file
        raise DocumentError(f"The Word document could not be read: {exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(parts)
