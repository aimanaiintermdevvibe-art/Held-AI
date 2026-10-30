import asyncio
import io
import json
import math
import re
import time
import zlib
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient

from backend.api.routers import agentic
from backend.ai_wrapper.contracts import AuditDecision, ClinicalAnalysis
from backend.ai_wrapper.orchestrator import AgenticPipeline
from backend.core.config import Settings
from backend.db import documents as store
from backend.main import app
from backend.rag import service
from backend.rag.chunking import chunk_text
from backend.ai_wrapper.client import ProviderBusyError
from backend.api.routers import documents as documents_router
from backend.rag.extract import DocumentError, extract_sections
from backend.rag.ingestion import IngestionWorker
from tests.test_agentic import FakeProvider

GUIDE = (
    "Willow House Visitor Guide\n\n"
    "Visiting hours are 10:00 to 19:30 every day. Please sign in at reception.\n\n"
    "The memory cafe meets every Thursday at 15:00 in the Garden Room. Families are welcome.\n\n"
    "Medication is given by the nursing team only; please do not bring medicine from home."
)


# -- fakes ---------------------------------------------------------------------

def bag_of_words(text: str) -> list[float]:
    """A deterministic stand-in for an embedding: texts sharing words point the same way."""
    vector = [0.0] * store.EMBEDDING_DIMENSIONS
    for word in re.findall(r"[a-z0-9]+", text.lower()):
        if len(word) > 3:
            vector[zlib.crc32(word.encode()) % store.EMBEDDING_DIMENSIONS] += 1.0
    vector[0] += 0.01  # never all zero
    return vector


def cosine(a, b):
    return sum(x * y for x, y in zip(a, b)) / (math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b)))


class FakeEmbedder:
    def __init__(self):
        self.calls: list[tuple[str, int]] = []

    async def embed(self, texts, *, model, task_type, dimensions):
        assert dimensions == store.EMBEDDING_DIMENSIONS
        self.calls.append((task_type, len(texts)))
        return [bag_of_words(t) for t in texts]


class FakeDocumentStore:
    """In-memory stand-in for backend.db.documents with the same owner scoping."""

    def __init__(self):
        self.documents: dict[str, dict] = {}
        self.chunks: dict[str, list[dict]] = {}

    def _view(self, doc):
        return {**doc, "chunk_count": sum(1 for c in self.chunks[doc["id"]] if c["embedding"] is not None)}

    async def find_document_by_hash(self, user_id, sha256):
        return next((self._view(d) for d in self.documents.values()
                     if d["user_id"] == user_id and d["sha256"] == sha256), None)

    async def create_document(self, *, user_id, filename, content_type, size_bytes, sha256, content,
                              embedding_model, metadata, chunks):
        doc_id = str(uuid4())
        self.documents[doc_id] = {
            "id": doc_id, "user_id": user_id, "filename": filename, "content_type": content_type,
            "size_bytes": size_bytes, "sha256": sha256, "status": "queued", "total_chunks": len(chunks),
            "error": None, "embedding_model": embedding_model, "metadata": metadata,
            "created_at": "2026-10-28T00:00:00+00:00",
        }
        self.chunks[doc_id] = [dict(c, chunk_index=i, id=f"{doc_id}:{i}", embedding=None) for i, c in enumerate(chunks)]
        return self._view(self.documents[doc_id]), True

    async def list_documents(self, user_id):
        return [self._view(d) for d in self.documents.values() if d["user_id"] == user_id]

    async def get_document(self, document_id, user_id=None, chunk_offset=0, chunk_limit=50):
        doc = self.documents.get(document_id)
        if not doc or (user_id is not None and doc["user_id"] != user_id):
            return None
        page = self.chunks[document_id][chunk_offset:chunk_offset + chunk_limit]
        return {**self._view(doc), "chunks": [
            {"chunk_index": c["chunk_index"], "content": c["content"], "metadata": c["metadata"],
             "embedded": c["embedding"] is not None} for c in page
        ]}

    async def delete_document(self, document_id, user_id):
        if not await self.get_document(document_id, user_id):
            return False
        del self.documents[document_id], self.chunks[document_id]
        return True

    async def has_documents(self, user_id):
        return any(c["embedding"] is not None for d in self.documents.values() if d["user_id"] == user_id
                   for c in self.chunks[d["id"]])

    async def unfinished_document_ids(self):
        return [d["id"] for d in self.documents.values() if d["status"] in ("queued", "processing", "paused")]

    async def set_document_status(self, document_id, status, error=None):
        self.documents[document_id].update(status=status, error=error)

    async def next_unembedded_chunks(self, document_id, limit):
        return [c for c in self.chunks[document_id] if c["embedding"] is None][:limit]

    async def save_chunk_embeddings(self, document_id, embeddings):
        by_id = {c["id"]: c for c in self.chunks[document_id]}
        for chunk_id, vector in embeddings:
            by_id[chunk_id]["embedding"] = vector
        return self._view(self.documents[document_id])["chunk_count"]

    async def search_chunks(self, user_id, embedding, limit, document_ids=None):
        found = [
            {"document_id": doc_id, "filename": self.documents[doc_id]["filename"],
             "chunk_index": c["chunk_index"], "metadata": c["metadata"], "content": c["content"],
             "similarity": round(cosine(embedding, c["embedding"]), 4)}
            for doc_id, chunks in self.chunks.items()
            if self.documents[doc_id]["user_id"] == user_id and (not document_ids or doc_id in document_ids)
            for c in chunks if c["embedding"] is not None
        ]
        return sorted(found, key=lambda r: r["similarity"], reverse=True)[:limit]


@pytest.fixture
def doc_store(monkeypatch):
    fake = FakeDocumentStore()
    for name in ("find_document_by_hash", "create_document", "list_documents", "get_document",
                 "delete_document", "has_documents", "search_chunks", "unfinished_document_ids",
                 "set_document_status", "next_unembedded_chunks", "save_chunk_embeddings"):
        monkeypatch.setattr(store, name, getattr(fake, name))
    return fake


@pytest.fixture(autouse=True)
def no_waiting(monkeypatch):
    """The worker's rate-limit and retry pauses return at once in tests."""
    real_sleep = asyncio.sleep

    async def quick(seconds, *args, **kwargs):
        await real_sleep(0)

    monkeypatch.setattr("backend.rag.ingestion.asyncio.sleep", quick)
    monkeypatch.setattr("backend.rag.ingestion.logger.propagate", True)  # lets caplog see progress lines


def ingest(settings, data=None, user_id="usr_1", filename="guide.txt", embedder=None):
    """Upload then run the worker to completion, as the API does in the background."""
    document, created = asyncio.run(service.ingest_document(
        settings, user_id=user_id, filename=filename, content_type="text/plain",
        data=GUIDE.encode() if data is None else data,
    ))
    if created or document["status"] == "queued":
        asyncio.run(IngestionWorker(embedder or FakeEmbedder(), settings).process(document["id"]))
    return document, created


@pytest.fixture
def embedder(monkeypatch):
    fake = FakeEmbedder()
    monkeypatch.setattr(agentic.gemini_client, "embed", fake.embed)
    return fake


@pytest.fixture
def client(fake_repo, auth, doc_store, embedder):
    with TestClient(app, headers=auth("usr_1")) as test_client:
        yield test_client


def upload(client, name="guide.txt", data=GUIDE.encode(), headers=None):
    return client.post("/agentic/documents", files={"file": (name, data, "text/plain")}, headers=headers)


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    objects = {1: b"<< /Type /Catalog /Pages 2 0 R >>", 3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    kids, number = [], 4
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects[number] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        objects[number + 1] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {number} 0 R >>"
        ).encode()
        kids.append(f"{number + 1} 0 R")
        number += 2
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for key in sorted(objects):
        offsets[key] = len(out)
        out += f"{key} 0 obj\n".encode() + objects[key] + b"\nendobj\n"
    xref, size = len(out), max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for key in range(1, size):
        out += f"{offsets[key]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- chunking ------------------------------------------------------------------

def test_chunks_respect_the_size_limit_and_overlap():
    text = "\n\n".join(" ".join(f"Sentence {p}.{s} has a few words." for s in range(10)) for p in range(8))
    chunks = chunk_text(text, max_chars=300, overlap=60)
    assert len(chunks) > 5
    assert all(len(c) <= 300 for c in chunks)
    for previous, current in zip(chunks, chunks[1:]):
        first_sentence = current.split(".")[0]
        assert first_sentence in previous  # each chunk starts with the end of the previous one


def test_short_text_is_one_chunk_and_unbroken_text_is_cut():
    assert chunk_text("Just one line.") == ["Just one line."]
    unbroken = chunk_text("x" * 2500, max_chars=1000, overlap=100)
    assert sum(map(len, unbroken)) == 2500 and max(map(len, unbroken)) <= 1000
    assert chunk_text("") == []


# -- extraction ----------------------------------------------------------------

def test_text_and_markdown_are_cleaned():
    [section] = extract_sections("notes.md", b"\xef\xbb\xbf# Title\r\n\r\n\r\n\r\nBody\x00 text   here")
    assert section.text == "# Title\n\nBody text here"
    assert section.page is None


def test_windows_encoded_text_is_decoded():
    [section] = extract_sections("old.txt", "Café opens at 9".encode("cp1252"))
    assert section.text == "Café opens at 9"


def test_pdf_pages_are_kept_as_sections():
    sections = extract_sections("guide.pdf", make_pdf(["Visiting hours start at ten", "", "The cafe is on page three"]))
    assert [(s.page, s.text) for s in sections] == [(1, "Visiting hours start at ten"), (3, "The cafe is on page three")]


def test_docx_paragraphs_and_tables_are_extracted():
    [section] = extract_sections("plan.docx", make_docx(["Morning routine", "Evening walk"], [["Day", "Time"], ["Mon", "10:00"]]))
    assert section.text == "Morning routine\n\nEvening walk\n\nDay | Time\n\nMon | 10:00"


@pytest.mark.parametrize("name, data, status", [
    ("image.png", b"\x89PNG", 415),
    ("noext", b"text", 415),
    ("blank.txt", b"   \n\n  ", 422),
    ("broken.pdf", b"not a pdf", 422),
    ("broken.docx", b"not a zip", 422),
])
def test_unusable_files_are_rejected(name, data, status):
    with pytest.raises(DocumentError) as caught:
        extract_sections(name, data)
    assert caught.value.status_code == status


# -- ingestion (upload + background worker) ------------------------------------

def test_upload_stores_chunks_and_the_worker_embeds_them_with_pages(doc_store):
    embed = FakeEmbedder()
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0)
    pdf = make_pdf(["Visiting hours are ten to seven", "The memory cafe meets on Thursdays"])
    document, created = asyncio.run(service.ingest_document(
        settings, user_id="usr_1", filename="C:\\fakepath\\guide.pdf", content_type="application/pdf", data=pdf,
    ))
    assert created
    assert (document["filename"], document["status"], document["chunk_count"], document["total_chunks"]) == ("guide.pdf", "queued", 0, 2)
    assert document["metadata"] == {"pages": 2}
    assert embed.calls == []  # nothing is embedded during the upload request

    assert asyncio.run(IngestionWorker(embed, settings).process(document["id"])) == "ready"
    assert embed.calls == [("RETRIEVAL_DOCUMENT", 2)]
    stored = doc_store.documents[document["id"]]
    assert stored["status"] == "ready"
    assert [c["metadata"] for c in doc_store.chunks[document["id"]]] == [{"page": 1}, {"page": 2}]
    assert all(c["embedding"] for c in doc_store.chunks[document["id"]])


def test_long_documents_have_no_limit_and_are_embedded_in_batches(doc_store, caplog):
    embed = FakeEmbedder()
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0, embedding_batch_size=100)
    text = "\n\n".join(f"Paragraph {i} talks about topic number {i} in some detail here." for i in range(800))
    with caplog.at_level("INFO", logger="held_ai.ingestion"):
        document, _ = ingest(settings, text.encode(), embedder=embed)
    total = doc_store.documents[document["id"]]["total_chunks"]
    assert total > 100 and settings.document_max_chunks == 0
    assert [n for _, n in embed.calls] == [100] * (total // 100) + ([total % 100] if total % 100 else [])
    progress = [r.getMessage() for r in caplog.records if "chunks embedded (" in r.getMessage()]
    assert len(progress) == len(embed.calls)
    assert f"{total}/{total} chunks embedded (100.0%)" in progress[-1]


def test_a_set_chunk_limit_still_rejects_before_storing(doc_store):
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0, document_max_chunks=2)
    with pytest.raises(DocumentError) as caught:
        asyncio.run(service.ingest_document(settings, user_id="usr_1", filename="long.txt",
                                            content_type=None, data=("word " * 400).encode()))
    assert caught.value.status_code == 413
    assert "split_document.py" in str(caught.value)
    assert doc_store.documents == {}


def test_same_file_is_stored_once_without_new_embeddings(doc_store):
    embed = FakeEmbedder()
    settings = Settings(_env_file=None)
    first, _ = ingest(settings, b"Hello there world", filename="a.txt", embedder=embed)
    second, created = ingest(settings, b"Hello there world", filename="b.txt", embedder=embed)
    assert (second["id"], created, second["status"]) == (first["id"], False, "ready")
    assert len(embed.calls) == 1


def test_rate_limits_are_waited_out(doc_store, caplog):
    class Busy(FakeEmbedder):
        async def embed(self, texts, **kwargs):
            if len(self.calls) < 2:
                self.calls.append(("busy", 0))
                raise ProviderBusyError(retry_after=4)
            return await super().embed(texts, **kwargs)

    with caplog.at_level("INFO", logger="held_ai.ingestion"):
        document, _ = ingest(Settings(_env_file=None), embedder=Busy())
    assert doc_store.documents[document["id"]]["status"] == "ready"
    assert sum("rate limited by Gemini" in r.getMessage() for r in caplog.records) == 2


def test_a_used_up_daily_quota_pauses_the_document_until_it_resets(doc_store, caplog):
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0, embedding_batch_size=1)
    long_text = "\n\n".join(f"Paragraph {i} talks about topic number {i} in some detail here." for i in range(40)).encode()

    class DailyQuota(FakeEmbedder):
        async def embed(self, texts, **kwargs):
            if len(self.calls) >= 3:
                raise ProviderBusyError(retry_after=72716)  # Gemini's answer once 1,000 a day are used
            return await super().embed(texts, **kwargs)

    with caplog.at_level("INFO", logger="held_ai.ingestion"):
        document, _ = ingest(settings, long_text, embedder=DailyQuota())
    stored = doc_store.documents[document["id"]]
    assert stored["status"] == "paused"
    assert "continues automatically" in stored["error"] and "3/" in stored["error"]
    assert any("paused at 3/" in r.getMessage() for r in caplog.records)
    assert asyncio.run(store.unfinished_document_ids()) == [document["id"]]  # resumed at startup too

    # When the quota resets, the worker continues from chunk 3.
    resumed = FakeEmbedder()
    assert asyncio.run(IngestionWorker(resumed, settings).process(document["id"])) == "ready"
    assert len(resumed.calls) == stored["total_chunks"] - 3


def test_a_failed_document_resumes_where_it_stopped_when_uploaded_again(doc_store):
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0, embedding_batch_size=1)

    class Breaks(FakeEmbedder):
        async def embed(self, texts, **kwargs):
            if len(self.calls) >= 2:
                raise RuntimeError("Gemini embedding request failed")
            return await super().embed(texts, **kwargs)

    long_text = "\n\n".join(f"Paragraph {i} talks about topic number {i} in some detail here." for i in range(40)).encode()
    document, _ = ingest(settings, long_text, embedder=Breaks())
    stored = doc_store.documents[document["id"]]
    assert stored["status"] == "failed"
    assert "upload the same file again to resume" in stored["error"]
    assert sum(c["embedding"] is not None for c in doc_store.chunks[document["id"]]) == 2

    resumed = FakeEmbedder()
    again, created = ingest(settings, long_text, embedder=resumed)
    assert (again["id"], created) == (document["id"], False)
    assert stored["status"] == "ready" and stored["error"] is None
    assert len(resumed.calls) == stored["total_chunks"] - 2  # only the chunks still missing


def test_unfinished_documents_are_resumed_at_startup(doc_store):
    settings = Settings(_env_file=None)
    document, _ = asyncio.run(service.ingest_document(settings, user_id="usr_1", filename="g.txt",
                                                      content_type=None, data=GUIDE.encode()))
    embed = FakeEmbedder()

    async def restart():
        worker = IngestionWorker(embed, settings)
        await worker.resume_unfinished()
        await worker.wait_until_idle()
        await worker.stop()

    asyncio.run(restart())
    assert doc_store.documents[document["id"]]["status"] == "ready"
    assert embed.calls


def test_an_instance_without_the_worker_leaves_documents_queued(doc_store):
    settings = Settings(_env_file=None, ingestion_worker=False)
    document, _ = asyncio.run(service.ingest_document(settings, user_id="usr_1", filename="g.txt",
                                                      content_type=None, data=GUIDE.encode()))
    embed = FakeEmbedder()

    async def start_and_upload():
        worker = IngestionWorker(embed, settings)
        await worker.resume_unfinished()
        worker.enqueue(document["id"])
        await worker.wait_until_idle()

    asyncio.run(start_and_upload())
    assert doc_store.documents[document["id"]]["status"] == "queued"
    assert embed.calls == []


def ingest_guide(doc_store, user_id="usr_1"):
    settings = Settings(_env_file=None, document_chunk_chars=200, document_chunk_overlap=0)
    return ingest(settings, user_id=user_id)[0]


def test_only_embedded_chunks_are_searched(doc_store):
    settings = Settings(_env_file=None)
    asyncio.run(service.ingest_document(settings, user_id="usr_1", filename="g.txt", content_type=None, data=GUIDE.encode()))
    assert asyncio.run(service.excerpts_for_chat(FakeEmbedder(), settings, user_id="usr_1", prompt="Summarise my document")) == []


def test_chat_retrieval_keeps_only_relevant_passages(doc_store):
    ingest_guide(doc_store)
    settings = Settings(_env_file=None, rag_min_similarity=0.2)
    results = asyncio.run(service.excerpts_for_chat(
        FakeEmbedder(), settings, user_id="usr_1", prompt="When does the memory cafe meet on Thursday?",
    ))
    assert len(results) == 1  # the medication passage shares no words and is left out
    assert "memory cafe" in results[0]["content"]
    assert results[0]["similarity"] >= 0.2


def test_unrelated_message_gets_no_passages_but_a_document_question_does(doc_store):
    ingest_guide(doc_store)
    settings = Settings(_env_file=None, rag_min_similarity=0.9)
    assert asyncio.run(service.excerpts_for_chat(FakeEmbedder(), settings, user_id="usr_1", prompt="I feel tired today")) == []
    summary = asyncio.run(service.excerpts_for_chat(FakeEmbedder(), settings, user_id="usr_1", prompt="Summarise my document"))
    assert len(summary) == settings.rag_top_k or len(summary) == len(doc_store.chunks[next(iter(doc_store.chunks))])


def test_no_embedding_call_for_users_without_documents(doc_store):
    embed = FakeEmbedder()
    assert asyncio.run(service.excerpts_for_chat(embed, Settings(_env_file=None), user_id="usr_1", prompt="hi")) == []
    assert embed.calls == []


def test_retrieval_failure_only_means_no_passages(doc_store, monkeypatch):
    ingest_guide(doc_store)

    class Down:
        async def embed(self, *args, **kwargs):
            raise RuntimeError("Gemini embedding request failed")

    assert asyncio.run(service.excerpts_for_chat(Down(), Settings(_env_file=None), user_id="usr_1", prompt="cafe")) == []


def test_excerpts_and_sources_carry_file_and_page():
    result = {"document_id": "d1", "filename": "guide.pdf", "chunk_index": 2, "metadata": {"page": 4},
              "content": "Text", "similarity": 0.81}
    assert service.as_excerpts([result]) == [{"source": "guide.pdf", "excerpt": "Text", "relevance": 0.81, "page": 4}]
    assert service.as_sources([result]) == [{"document_id": "d1", "filename": "guide.pdf", "chunk_index": 2,
                                             "page": 4, "similarity": 0.81}]


# -- endpoints -----------------------------------------------------------------

def wait_until_ready(client, document_id, timeout=5.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        document = client.get(f"/agentic/documents/{document_id}").json()
        if document["status"] in ("ready", "failed"):
            return document
        time.sleep(0.02)
    raise AssertionError(f"document still {document['status']}")


def test_upload_is_accepted_then_embedded_in_the_background(client, embedder, doc_store):
    res = upload(client)
    assert res.status_code == 202
    document = res.json()["document"]
    assert res.json()["created"] is True
    assert document["status"] in ("queued", "processing", "ready") and document["total_chunks"] >= 1

    ready = wait_until_ready(client, document["id"])
    assert (ready["status"], ready["chunk_count"]) == ("ready", ready["total_chunks"])
    assert all(c["embedded"] for c in ready["chunks"])
    assert ("RETRIEVAL_DOCUMENT", ready["total_chunks"]) in embedder.calls

    again = upload(client)
    assert (again.status_code, again.json()["created"], again.json()["document"]["id"]) == (200, False, document["id"])

    assert [d["id"] for d in client.get("/agentic/documents").json()] == [document["id"]]
    assert "memory cafe" in " ".join(c["content"] for c in ready["chunks"])

    found = client.post("/agentic/documents/search", json={"query": "memory cafe Thursday", "top_k": 2}).json()
    assert found["results"][0]["document_id"] == document["id"]
    assert "memory cafe" in found["results"][0]["content"]
    assert embedder.calls[-1] == ("RETRIEVAL_QUERY", 1)

    assert client.delete(f"/agentic/documents/{document['id']}").json() == {"success": True}
    assert client.get(f"/agentic/documents/{document['id']}").status_code == 404
    assert client.get("/agentic/documents").json() == []


def test_chunks_are_listed_a_page_at_a_time(client, monkeypatch):
    monkeypatch.setattr(documents_router.settings, "document_chunk_chars", 200)
    monkeypatch.setattr(documents_router.settings, "document_chunk_overlap", 0)
    text = "\n\n".join(f"Paragraph {i} talks about topic number {i} in some detail here." for i in range(40))
    document = upload(client, "long.txt", text.encode()).json()["document"]
    page = client.get(f"/agentic/documents/{document['id']}?chunk_offset=2&chunk_limit=3").json()
    assert [c["chunk_index"] for c in page["chunks"]] == [2, 3, 4]
    assert page["total_chunks"] > 5


def test_documents_are_private_to_their_owner(client, auth):
    document = upload(client).json()["document"]
    wait_until_ready(client, document["id"])
    other = auth("usr_2")
    assert client.get("/agentic/documents", headers=other).json() == []
    assert client.get(f"/agentic/documents/{document['id']}", headers=other).status_code == 404
    assert client.delete(f"/agentic/documents/{document['id']}", headers=other).status_code == 404
    assert client.post("/agentic/documents/search", json={"query": "cafe"}, headers=other).json()["results"] == []


@pytest.mark.parametrize("name, data, status", [("photo.png", b"\x89PNG", 415), ("empty.txt", b"", 422)])
def test_unusable_uploads_get_a_clear_error(client, name, data, status):
    res = upload(client, name, data)
    assert res.status_code == status
    assert res.json()["detail"]


def test_upload_requires_a_signed_in_user(fake_repo, doc_store, embedder):
    with TestClient(app) as anonymous:
        assert anonymous.post("/agentic/documents", files={"file": ("a.txt", b"hi", "text/plain")}).status_code == 401


def test_embedding_outage_marks_the_document_failed_not_the_upload(client, monkeypatch):
    async def down(*args, **kwargs):
        raise RuntimeError("Gemini embedding request failed")

    monkeypatch.setattr(agentic.gemini_client, "embed", down)
    res = upload(client)
    assert res.status_code == 202
    failed = wait_until_ready(client, res.json()["document"]["id"])
    assert failed["status"] == "failed" and "resume" in failed["error"]


# -- chat (RAG) ----------------------------------------------------------------

@pytest.fixture
def rag_chat(monkeypatch, client, doc_store, embedder):
    provider = FakeProvider(route="CLINICAL_QUERY")
    monkeypatch.setattr(agentic, "pipeline", AgenticPipeline(provider, Settings(_env_file=None)))
    monkeypatch.setattr(agentic, "excerpts_for_chat", service.excerpts_for_chat)  # undo the conftest stub
    monkeypatch.setattr(agentic.settings, "rag_min_similarity", 0.15)  # the bag-of-words fake scores low
    monkeypatch.setattr(agentic.settings, "agentic_stream_delay_seconds", 0.0)

    async def title(*args):
        return "Title"

    monkeypatch.setattr(agentic, "generate_title", title)
    wait_until_ready(client, upload(client).json()["document"]["id"])
    return client, provider


def test_chat_reply_is_given_the_relevant_passages_and_returns_sources(rag_chat):
    client, provider = rag_chat
    res = client.post("/agentic/run", json={"prompt": "When is the memory cafe on Thursday?"})
    assert res.status_code == 200
    sources = res.json()["sources"]
    assert sources and sources[0]["filename"] == "guide.txt"

    excerpts = provider.payloads[-1]["document_excerpts"]
    assert excerpts[0]["source"] == "guide.txt" and "memory cafe" in excerpts[0]["excerpt"]
    inputs = {model: json.loads(text) for model, text in provider.structured_inputs if model in (ClinicalAnalysis, AuditDecision)}
    assert inputs[ClinicalAnalysis]["document_excerpts"] == excerpts
    assert inputs[AuditDecision]["document_excerpts"] == excerpts


def test_streamed_reply_reports_its_sources_when_done(rag_chat):
    client, _ = rag_chat
    res = client.post("/agentic/stream", json={"prompt": "When is the memory cafe on Thursday?"})
    events = [json.loads(line[6:]) for line in res.text.splitlines() if line.startswith("data:")]
    assert events[-1]["done"] is True
    assert events[-1]["sources"][0]["filename"] == "guide.txt"


def test_off_topic_questions_still_use_the_documents(rag_chat):
    client, provider = rag_chat
    provider.intent = "OFF_TOPIC"
    client.post("/agentic/run", json={"prompt": "What are the visiting hours at reception?"})
    assert provider.payloads[-1]["document_excerpts"]


def test_unrelated_chat_gets_no_passages(rag_chat, monkeypatch):
    client, provider = rag_chat
    monkeypatch.setattr(agentic.settings, "rag_min_similarity", 0.99)
    res = client.post("/agentic/run", json={"prompt": "I feel lonely tonight."})
    assert "sources" not in res.json()
    assert "document_excerpts" not in provider.payloads[-1]
